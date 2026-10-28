#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build-texts-pack.py: TEXTS PACK of the Lepa Svaki Dan v2 previews

Reads the 21 Word documents (the treatment texts Tamara wrote with a copywriter) from
<project>/procedure/<4 segment folders>/*.docx with python-docx and writes, next to content.json:

    articles.json / articles.js      21 articles: blocks (h2|h3|p|ul|ol), excerpt, readMinutes, placeholders, links
    segments.json                    4 segments: ordered articles, price-list tabs/groups, hero image, intro quote
    service-links.json               article <-> price-list groups/services + reverse index
    search-synonyms.json             HIDDEN search vocabulary (never shown as hashtags)
    CHANGES.md                       every replacement made by the content fixes
    MANIFEST-ARTICLES.md             what every file is, counts, article table, verification

The same reviewed fix rules as content.json are applied (content-fixes.rules.json): the owner is not a
doctor, no em dash, en dash only in numeric ranges. Then verify-texts-pack.py runs and its report is embedded.

Usage:  python3 _tools/build-texts-pack.py        (needs python-docx and lxml)
ECHOSYSTEM-MOTHER AGENCY
"""
import collections
import datetime
import importlib.util
import io
import json
import math
import os
import re
import sys
import unicodedata
import zipfile

sys.dont_write_bytecode = True
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import lsd_common as C            # noqa: E402
import texts_curation as K        # noqa: E402
import texts_search_data as SD    # noqa: E402

from docx import Document                 # noqa: E402
from docx.oxml.ns import qn               # noqa: E402

NOW = datetime.datetime.now(datetime.timezone.utc)
TODAY = NOW.strftime('%Y-%m-%d')
EM, EN = C.EM, C.EN
NBSP = chr(0xA0)


def nfc(s):
    return unicodedata.normalize('NFC', s)


# ============================================================================ 1. read the Word documents
def numbering_formats(doc):
    """numId -> numFmt of level 0 ('bullet', 'decimal', ...)."""
    try:
        root = doc.part.numbering_part.element
    except Exception:
        return {}
    abstract = {}
    for a in root.findall(qn('w:abstractNum')):
        lvl0 = [l for l in a.findall(qn('w:lvl')) if l.get(qn('w:ilvl')) == '0']
        if lvl0:
            abstract[a.get(qn('w:abstractNumId'))] = lvl0[0].find(qn('w:numFmt')).get(qn('w:val'))
    out = {}
    for n in root.findall(qn('w:num')):
        out[int(n.get(qn('w:numId')))] = abstract.get(n.find(qn('w:abstractNumId')).get(qn('w:val')))
    return out


def raw_text(el):
    """Text of a w:p straight from the XML (independent of python-docx's own text property)."""
    parts = []
    for node in el.iter():
        if node.tag == qn('w:t'):
            parts.append(node.text or '')
        elif node.tag == qn('w:tab'):
            parts.append('\t')
        elif node.tag in (qn('w:br'), qn('w:cr')):
            parts.append('\n')
        elif node.tag == qn('w:noBreakHyphen'):
            parts.append('-')
    return ''.join(parts)


def read_comments(path, doc, para_index):
    """[{id, text, anchorText, paraIndex}] from word/comments.xml and the comment ranges in document order."""
    z = zipfile.ZipFile(path)
    if 'word/comments.xml' not in z.namelist():
        return []
    from lxml import etree
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    croot = etree.fromstring(z.read('word/comments.xml'))
    texts = {}
    for c in croot.findall(W + 'comment'):
        texts[c.get(W + 'id')] = '\n'.join(''.join(t.text or '' for t in p.iter(W + 't')) for p in c.iter(W + 'p')).strip()
    droot = doc.element.body
    active, found, cur = {}, {}, None
    for node in droot.iter():
        if node.tag == W + 'p':
            cur = para_index.get(node, cur)
        elif node.tag == W + 'commentRangeStart':
            active[node.get(W + 'id')] = {'para': cur, 'buf': []}
        elif node.tag == W + 'commentRangeEnd':
            cid = node.get(W + 'id')
            if cid in active:
                found[cid] = active.pop(cid)
        elif node.tag == W + 't':
            for a in active.values():
                a['buf'].append(node.text or '')
    out = []
    for cid in sorted(texts, key=int):
        f = found.get(cid, {'para': None, 'buf': []})
        out.append({'id': int(cid), 'text': texts[cid], 'anchorText': ''.join(f['buf']).strip(), 'paraIndex': f['para']})
    return out


def parse_docx(path):
    doc = Document(path)
    fmts = numbering_formats(doc)
    paras = []
    mismatch = 0
    para_index = {}
    for i, p in enumerate(doc.paragraphs):
        para_index[p._p] = i
        ppr = p._p.pPr
        num = ppr.numPr if ppr is not None else None
        num_id = num.numId.val if (num is not None and num.numId is not None) else None
        links = []
        for h in p._p.iter(qn('w:hyperlink')):
            rid = h.get(qn('r:id'))
            href = doc.part.rels[rid].target_ref if rid in doc.part.rels else None
            links.append({'text': ''.join(t.text or '' for t in h.iter(qn('w:t'))), 'href': href})
        runs = [r for r in p.runs if r.text.strip()]
        all_bold = bool(runs) and all(bool(r.bold) for r in runs)
        txt = p.text
        if txt != raw_text(p._p):
            mismatch += 1
        paras.append({'idx': i, 'style': p.style.name if p.style is not None else '', 'text': txt, 'numId': num_id,
                      'numFmt': fmts.get(num_id) if num_id is not None else None, 'links': links, 'allBold': all_bold})
    comments = read_comments(path, doc, para_index)
    return {'paras': paras, 'comments': comments, 'textMismatch': mismatch}


def list_documents():
    docs = []
    for seg in K.SEGMENTS:
        dirs = [d for d in os.listdir(C.PROCEDURE_DIR) if nfc(d) == seg['folder']]
        if len(dirs) != 1:
            raise SystemExit('segment folder not found on disk: %r (found %r)' % (seg['folder'], dirs))
        folder = dirs[0]
        files = sorted([f for f in os.listdir(os.path.join(C.PROCEDURE_DIR, folder)) if f.lower().endswith('.docx') and not f.startswith('~')], key=nfc)
        for f in files:
            docs.append({'segment': seg, 'folder': folder, 'file': f, 'path': os.path.join(C.PROCEDURE_DIR, folder, f),
                         'sourceFile': nfc('procedure/' + folder + '/' + f)})
    return docs


# ============================================================================ 2. text processing
BLANK_RE = re.compile(r'_{3,}')
PLACEHOLDER_RE = re.compile(r'\[\[dopuniti: ([^\]]+)\]\]')


def placeholder_kind(before, after):
    """Chosen from the sentence context, never guessed: 'dr ____' = a doctor, '____ laser' = a laser device."""
    if re.search(r'\bdr\.?\s*$', before, re.I):
        return 'ime lekara'
    if re.match(r'\s*laser', after, re.I):
        return 'naziv lasera'
    raise SystemExit('cannot decide what the blank is from its context: ...%s[____]%s...' % (before[-40:], after[:40]))


def convert_blanks(t):
    """'____' -> [[dopuniti: ...]] (kind chosen from the sentence context)."""
    out, pos = [], 0
    for m in BLANK_RE.finditer(t):
        out.append(t[pos:m.start()])
        out.append('[[dopuniti: %s]]' % placeholder_kind(t[:m.start()], t[m.end():]))
        pos = m.end()
    out.append(t[pos:])
    return ''.join(out)


class FixLog:
    def __init__(self):
        self.items = []

    def add(self, file, where, before, after, rules):
        self.items.append({'file': file, 'where': where, 'before': before, 'after': after, 'rules': rules})


def process_text(raw, file, where, fixlog, stats):
    """whitespace -> placeholders -> reviewed content fixes. Returns the final text."""
    t = raw
    n_nbsp = t.count(NBSP)
    if n_nbsp:
        stats['nbsp'] += n_nbsp
        t = t.replace(NBSP, ' ')
    t = re.sub(r'[ \t]+', ' ', t).strip()
    if '\n' in t:
        raise SystemExit('line break inside a paragraph in %s (%s): not expected' % (file, where))
    t = convert_blanks(t)
    if '_' in t:
        raise SystemExit('stray underscore left in %s (%s): %r' % (file, where, t))
    hits = []
    fixed = C.fix_text(t, hits)
    if fixed != t:
        fixlog.add(file, where, t, fixed, hits)
    C.assert_clean(fixed, '%s (%s)' % (file, where))
    return fixed


def split_title(first_line):
    """title = part before the FIRST colon, subtitle = the rest (empty when there is no colon)."""
    t = re.sub(r'[ \t]+', ' ', first_line.replace(NBSP, ' ')).strip()
    if ':' in t:
        a, b = t.split(':', 1)
        return a.strip(), b.strip()
    return t, ''


def make_excerpt(text, limit=170):
    t = text.strip()
    if len(t) <= limit:
        out = t
    else:
        out = t[:limit] if t[limit] == ' ' else (t[:t[:limit + 1].rfind(' ')] if ' ' in t[:limit + 1] else t[:limit])
    out = out.rstrip(' ,;:-(')
    # never leave an opened bracket or quote dangling
    for o, c in (('(', ')'), (chr(0x201E), chr(0x201C))):
        while out.count(o) > out.count(c):
            out = out[:out.rfind(o)].rstrip(' ,;:-')
    return out


# ============================================================================ 3. build articles
def build_articles(docs, content, fixlog, report):
    stats = {'nbsp': 0, 'pseudoHeadingCandidates': [], 'paragraphTextMismatch': 0, 'emptyParagraphs': 0}
    blog = content['blog']
    articles = []
    for n, d in enumerate(docs, 1):
        parsed = parse_docx(d['path'])
        stats['paragraphTextMismatch'] += parsed['textMismatch']
        paras = parsed['paras']
        nonempty = [p for p in paras if p['text'].strip()]
        stats['emptyParagraphs'] += len(paras) - len(nonempty)
        title_para = nonempty[0]
        title_raw, subtitle_raw = split_title(title_para['text'])
        title = process_text(title_raw, d['sourceFile'], 'title', fixlog, stats)
        subtitle = process_text(subtitle_raw, d['sourceFile'], 'subtitle', fixlog, stats) if subtitle_raw else ''
        blocks, para2block, placeholders = [], {}, []
        prev_list = None  # (numId, block index) of the previous non-empty paragraph if it was a list item
        for p in nonempty[1:]:
            where_idx = len(blocks)
            style = p['style']
            if style in ('Heading 1', 'Heading 2', 'Heading 3'):
                if style == 'Heading 1':
                    raise SystemExit('second Heading 1 in %s' % d['file'])
                blocks.append({'type': 'h2' if style == 'Heading 2' else 'h3',
                               'text': process_text(p['text'], d['sourceFile'], 'block %d' % where_idx, fixlog, stats)})
                prev_list = None
            elif p['numId'] is not None:
                kind = 'ol' if p['numFmt'] == 'decimal' else 'ul'
                if prev_list is not None and prev_list[0] == p['numId'] and blocks and blocks[-1]['type'] == kind:
                    bi = len(blocks) - 1
                else:
                    blocks.append({'type': kind, 'items': []})
                    bi = len(blocks) - 1
                item_no = len(blocks[bi]['items'])
                blocks[bi]['items'].append(process_text(p['text'], d['sourceFile'], 'block %d item %d' % (bi, item_no), fixlog, stats))
                prev_list = (p['numId'], bi)
                where_idx = bi
            else:
                raw = re.sub(r'[ \t]+', ' ', p['text']).strip()
                is_q = p['allBold'] and len(raw) <= 140 and raw.endswith('?') and raw.count('?') == 1 and '.' not in raw
                if is_q:
                    stats['pseudoHeadingCandidates'].append('%s: %s' % (d['file'], raw))
                    blocks.append({'type': 'h2', 'text': process_text(p['text'], d['sourceFile'], 'block %d' % where_idx, fixlog, stats)})
                else:
                    blocks.append({'type': 'p', 'text': process_text(p['text'], d['sourceFile'], 'block %d' % where_idx, fixlog, stats)})
                prev_list = None
            para2block[p['idx']] = (len(blocks) - 1, len(blocks[-1]['items']) - 1 if 'items' in blocks[-1] else None)
            if p['links']:
                bi = len(blocks) - 1
                blk = blocks[bi]
                lk = blk.setdefault('links', [])
                for l in p['links']:
                    entry = {'text': C.fix_text(re.sub(r'[ \t]+', ' ', l['text'].replace(NBSP, ' ')).strip()), 'legacyHref': l['href']}
                    if blk['type'] in ('ul', 'ol'):
                        entry = {'item': len(blk['items']) - 1, **entry}
                    lk.append(entry)
        # placeholders (positions in the FINAL text)
        for bi, b in enumerate(blocks):
            texts = [(None, b['text'])] if 'text' in b else list(enumerate(b['items']))
            for item, t in texts:
                for m in PLACEHOLDER_RE.finditer(t):
                    placeholders.append({'token': m.group(0), 'kind': m.group(1), 'blockIndex': bi, 'blockType': b['type'],
                                         **({'itemIndex': item} if item is not None else {}),
                                         'context': t[max(0, m.start() - 45):m.end() + 25].strip()})
        # flag the one paragraph that is an editorial source note (not article copy)
        for b in blocks:
            if b['type'] == 'p' and b['text'].startswith('Informacije su preuzete sa ovog sajta'):
                b['flag'] = 'source-note: copywriter note naming the site the text was taken from; not article copy, do not publish'
        # excerpt, counts
        first_p = next(b for b in blocks if b['type'] == 'p')
        excerpt = make_excerpt(first_p['text'])
        assert '[[' not in excerpt, 'placeholder inside excerpt of ' + d['file']
        body_words = 0
        for b in blocks:
            body_words += C.word_count(b['text']) if 'text' in b else sum(C.word_count(i) for i in b['items'])
        title_words = C.word_count(title) + (C.word_count(subtitle) if subtitle else 0)
        full_title = title + (': ' + subtitle if subtitle else '')
        slug = C.slugify(title)
        # legacy blog post (old blog.html teaser data in content.json)
        i0 = n - 1
        lb = blog[i0]
        match = 'exact' if nfc(lb['title']) == nfc(full_title) else 'order'
        art = {
            'id': 'art-%d' % n, 'slug': slug, 'segmentId': d['segment']['id'], 'title': title, 'subtitle': subtitle,
            'blocks': blocks, 'excerpt': excerpt, 'readMinutes': max(1, math.ceil(body_words / 200)),
            'sourceFile': d['sourceFile'], 'hasPlaceholders': bool(placeholders), 'placeholders': placeholders,
            'wordCount': body_words, 'titleWordCount': title_words,
            'legacy': {'blogIndex': i0, 'blogTitle': lb['title'], 'titleMatch': match, 'blogWordCount': lb['wordCount'], 'blogHadBlanks': lb['hasPlaceholderBlanks']},
            '_comments': parsed['comments'], '_para2block': para2block, '_key': None,
        }
        articles.append(art)
    # keys + slug check
    assert len(articles) == len(K.ARTICLE_KEYS), (len(articles), len(K.ARTICLE_KEYS))
    for a, (key, slug) in zip(articles, K.ARTICLE_KEYS):
        if a['slug'] != slug:
            raise SystemExit('article %s has slug %r, curation expects %r for key %r (order or title changed: re-check texts_curation.py)' % (a['id'], a['slug'], slug, key))
        a['_key'] = key
    slugs = [a['slug'] for a in articles]
    assert len(set(slugs)) == len(slugs), 'duplicate slugs'
    report['docStats'] = stats
    return articles


def resolve_url(href, key_to_id):
    if not href:
        return {'type': 'none'}
    m = re.match(r'https?://(?:www\.)?lepasvakidan\.rs(/[^?#]*)?', href)
    if not m:
        return {'type': 'external', 'note': 'source site cited by the copywriter'} if 'preimedermafacial' in href else {'type': 'external'}
    path = m.group(1) or '/'
    if not path.endswith('/'):
        path += '/'
    t = K.URL_TARGETS.get(path)
    if not t:
        return {'type': 'none', 'note': 'unmapped legacy path ' + path}
    out = {'type': t['type']}
    if t['type'] == 'article':
        out['id'] = key_to_id[t['key']]
    elif 'id' in t:
        out['id'] = t['id']
    if 'note' in t:
        out['note'] = t['note']
    return out


def finish_articles(articles, key_to_id, seg_names):
    """cross-link hints from the Word comments, link targets; then strip the private fields."""
    n_links, n_hints = 0, 0
    for a in articles:
        for b in a['blocks']:
            for l in b.get('links', []):
                l['target'] = resolve_url(l['legacyHref'], key_to_id)
                n_links += 1
        hints = []
        for c in a['_comments']:
            res = K.COMMENT_RESOLUTION.get((a['_key'], c['id']))
            if res is None:
                raise SystemExit('Word comment %s #%d of %s has no resolution in texts_curation.COMMENT_RESOLUTION' % (a['id'], c['id'], a['sourceFile']))
            bi = None
            if c['paraIndex'] is not None and c['paraIndex'] in a['_para2block']:
                bi = a['_para2block'][c['paraIndex']][0]
            anchor = C.fix_text(convert_blanks(re.sub(r'[ \t]+', ' ', c['anchorText'].replace(NBSP, ' ')).strip()))
            hints.append({
                'comment': C.fix_text(c['text']), 'anchorText': anchor, 'blockIndex': bi,
                'resolution': res['resolution'], 'basis': res['basis'],
                'targetArticleIds': [key_to_id[k] for k in res.get('articles', [])],
                'targetSegmentIds': list(res.get('segments', [])),
            })
            n_hints += 1
        a['crossLinkHints'] = hints
    out = []
    for a in articles:
        d = {k: v for k, v in a.items() if k == '_key' or not k.startswith('_')}
        out.append(d)
    return out, n_links, n_hints


# ============================================================================ 4. price-list links
def tab_order(content):
    return {t['id']: i for i, t in enumerate(content['pricing']['tabs'])}


def sid_sort_key(order):
    def k(sid):
        t, g, s = sid.split(':')
        return (order[t], int(g), int(s))
    return k


def build_service_links(articles, content, key_to_id):
    tabs = {t['id']: t for t in content['pricing']['tabs']}
    order = tab_order(content)
    svc_by_id = {s['id']: s for s in content['allServices']}
    fold = C.fold
    per_article = {}
    for a in articles:
        key = [k for k, i in key_to_id.items() if i == a['id']][0]
        specs = K.LINKS[key]
        groups = collections.OrderedDict()   # groupId -> link dict
        primary_taken = set()

        def group_of(gid):
            t, gi = gid.split(':')
            return tabs[t], int(gi), tabs[t]['groups'][int(gi)]

        def add_services(link, ids, basis, why, origin=None):
            have = {s['id'] for s in link['services']}
            for sid in ids:
                if sid not in svc_by_id:
                    raise SystemExit('unknown service id %s in LINKS[%s]' % (sid, key))
                if sid in have:
                    continue
                link['services'].append({'id': sid, 'name': svc_by_id[sid]['name'], 'basis': basis, 'why': why, **({'origin': origin} if origin else {})})

        for spec in specs:
            if spec['kind'] == 'group':
                tab, gi, grp = group_of(spec['group'])
                link = groups.setdefault(spec['group'], {
                    'tabId': tab['id'], 'groupIndex': gi, 'groupId': spec['group'], 'groupTitle': grp['title'], 'tabLabel': tab['label'],
                    'role': spec['role'], 'basis': spec['basis'], 'origin': spec['origin'], 'services': []})
                for sel in spec['sels']:
                    ids = []
                    sv = grp['services']
                    for s in sv:
                        sid = s['id']
                        if sel.get('all'):
                            ids.append(sid)
                        elif 'idx' in sel and int(sid.split(':')[2]) in sel['idx']:
                            ids.append(sid)
                        elif 'ids' in sel and sid in sel['ids']:
                            ids.append(sid)
                        elif 'name' in sel and re.search(sel['name'], fold(s['name'])):
                            ids.append(sid)
                        elif 'namedesc' in sel and re.search(sel['namedesc'], fold(s['name'] + ' ' + (s.get('description') or ''))):
                            ids.append(sid)
                    for sid in sel.get('ids', []):
                        if not sid.startswith(spec['group'] + ':'):
                            raise SystemExit('LINKS[%s]: service %s is not inside group %s' % (key, sid, spec['group']))
                    if not ids:
                        raise SystemExit('LINKS[%s] group %s: selector %r matched no service' % (key, spec['group'], sel))
                    add_services(link, ids, sel['basis'], sel['why'], spec['origin'])
            else:
                rx = re.compile(spec['pattern'])
                found = collections.OrderedDict()
                for tid in spec['tabs']:
                    for gi, grp in enumerate(tabs[tid]['groups']):
                        gid = '%s:%d' % (tid, gi)
                        if gid in spec['skip'] or (spec['only_groups'] and gid not in spec['only_groups']):
                            continue
                        for s in grp['services']:
                            if rx.search(fold(s['name'] + ' ' + (s.get('description') or ''))):
                                found.setdefault(gid, []).append(s['id'])
                if not found:
                    raise SystemExit('LINKS[%s]: scan %r matched nothing' % (key, spec['pattern']))
                for gid, ids in found.items():
                    if gid in groups:
                        continue
                    tab, gi, grp = group_of(gid)
                    link = groups.setdefault(gid, {'tabId': tab['id'], 'groupIndex': gi, 'groupId': gid, 'groupTitle': grp['title'], 'tabLabel': tab['label'],
                                                   'role': spec['role'], 'basis': spec['basis'], 'origin': spec['origin'], 'services': []})
                    add_services(link, ids, 'name or description', 'name or description contains the topic of the article')
        links = list(groups.values())
        for l in links:
            l['services'].sort(key=lambda s: sid_sort_key(order)(s['id']))
            l['serviceIds'] = [s['id'] for s in l['services']]
        links.sort(key=lambda l: (0 if l['role'] == 'primary' else 1, order[l['tabId']], l['groupIndex']))
        per_article[a['id']] = links
    return per_article


def reconcile_old_rules(content, per_article, key_to_id):
    out = []
    for r in content['serviceBlogLinks']:
        art_id = 'art-%d' % (r['postIndex'] + 1)
        got = set()
        for l in per_article[art_id]:
            got.update(l['serviceIds'])
        missing = [s for s in r['linkedServiceIds'] if s not in got]
        if missing:
            raise SystemExit('old SERVICE_BLOG_LINKS rule %r: services %s are not linked to %s' % (r['matchSubstring'], missing, art_id))
        out.append({'oldRule': 'SERVICE_BLOG_LINKS', 'matchSubstring': r['matchSubstring'], 'postIndex': r['postIndex'], 'postTitle': r['postTitle'],
                    'articleId': art_id, 'serviceIds': r['linkedServiceIds'], 'reconciled': True,
                    'note': 'every service of the old rule is linked to the same text here (the old post index equals the article order)'})
    return out


def render_service_links(articles, content, per_article, old_rules):
    order = tab_order(content)
    tabs = content['pricing']['tabs']
    all_groups = {}
    for t in tabs:
        for gi, g in enumerate(t['groups']):
            all_groups['%s:%d' % (t['id'], gi)] = {'tabId': t['id'], 'tabLabel': t['label'], 'groupIndex': gi, 'title': g['title'],
                                                  'serviceCount': len(g['services']), 'primaryArticleIds': [], 'relatedArticleIds': []}
    svc_rev = collections.OrderedDict()
    arts = []
    unmatched = []
    for a in articles:
        links = per_article[a['id']]
        if not links:
            unmatched.append(a['id'])
        for l in links:
            gi = all_groups[l['groupId']]
            (gi['primaryArticleIds'] if l['role'] == 'primary' else gi['relatedArticleIds']).append(a['id'])
            for sid in l['serviceIds']:
                svc_rev.setdefault(sid, [])
                if a['id'] not in svc_rev[sid]:
                    svc_rev[sid].append(a['id'])
        arts.append({
            'articleId': a['id'], 'slug': a['slug'], 'title': a['title'], 'segmentId': a['segmentId'], 'matched': bool(links),
            'primaryGroupIds': [l['groupId'] for l in links if l['role'] == 'primary'],
            'relatedGroupIds': [l['groupId'] for l in links if l['role'] == 'related'],
            'serviceIds': sorted({s for l in links for s in l['serviceIds']}, key=sid_sort_key(order)),
            'groups': [{k: v for k, v in l.items() if k != 'services'} | {'services': l['services']} for l in links],
        })
    for g in all_groups.values():
        g['articleIds'] = g['primaryArticleIds'] + [x for x in g['relatedArticleIds'] if x not in g['primaryArticleIds']]
    without = [gid for gid, g in all_groups.items() if not g['articleIds']]
    svc_rev_sorted = collections.OrderedDict(sorted(svc_rev.items(), key=lambda kv: sid_sort_key(order)(kv[0])))
    return {
        'meta': {
            'pack': 'Lepa Svaki Dan v2 texts pack: service-links.json', 'generatedBy': 'ECHOSYSTEM-MOTHER AGENCY', 'generatedAt': NOW.strftime('%Y-%m-%dT%H:%M:%SZ'),
            'ids': {'groupId': '"<tabId>:<groupIndex>" (the data-group-id convention of the previews)', 'serviceId': '"<tabId>:<groupIndex>:<serviceIndex>" (content.json allServices[].id)',
                    'articleId': 'art-1 ... art-21 (articles.json)'},
            'roles': {'primary': 'the price-list group IS the subject of the article (or holds the services named in it)',
                      'related': 'other groups whose services are named in, or combined with, the article subject'},
            'basis': {'group': 'whole group is the subject', 'name': 'service name contains the topic', 'description': 'service description contains the topic',
                      'name or description': 'name or description contains the topic', 'text': 'the article text names it, the service name does not'},
            'method': 'Curated in _tools/texts_curation.py and VALIDATED by the builder: every id exists, every selector matches at least one service, the two old SERVICE_BLOG_LINKS rules are reproduced. Matching is on the real service name/description (diacritics folded), never on invented synonyms.',
            'counts': {'articles': len(arts), 'matchedArticles': len(arts) - len(unmatched), 'unmatchedArticles': len(unmatched), 'groups': len(all_groups),
                       'groupsWithoutArticle': len(without), 'servicesLinked': len(svc_rev_sorted), 'services': len(content['allServices'])},
        },
        'oldSiteRules': old_rules,
        'unmatchedArticles': [{'articleId': u} for u in unmatched],
        'unmatchedNote': 'Every one of the 21 articles has at least one price-list group.' if not unmatched else 'Articles without any price-list group are listed in unmatchedArticles.',
        'articles': arts,
        'groups': all_groups,
        'groupsWithoutArticle': [{'groupId': g, 'title': all_groups[g]['title'], 'tabId': all_groups[g]['tabId']} for g in without],
        'serviceToArticles': svc_rev_sorted,
    }


# ============================================================================ 5. segments
def sentence_of(text, n):
    parts = re.split(r'(?<=[.!?])\s+(?=[A-ZČĆŠŽĐ])', text.strip())
    return parts[n - 1]


def get_path(obj, path):
    cur = obj
    for tok in re.findall(r'[^.\[\]]+', path):
        cur = cur[int(tok)] if isinstance(cur, list) else cur[tok]
    return cur


def resolve_quote(q, articles_by_key, content):
    if q['kind'] == 'article':
        a = articles_by_key[q['article']]
        b = a['blocks'][q['block']]
        s = sentence_of(b['text'], q['sentence'])
        assert s in b['text']
        return {'text': s, 'source': {'kind': 'article', 'articleId': a['id'], 'blockIndex': q['block'], 'sentence': q['sentence']}}
    s = get_path(content, q['path'])
    return {'text': s, 'source': {'kind': 'content.json', 'path': q['path'], 'note': 'category line of the old site'}}


def build_segments(articles, content, per_article, img_manifest):
    by_key = {a['_key']: a for a in articles}
    imgs = {im['id']: im for im in img_manifest['images']}
    order = tab_order(content)
    segs = []
    for sd in K.SEGMENTS:
        arts = [a for a in articles if a['segmentId'] == sd['id']]
        # price-list relation
        rel = collections.OrderedDict()
        for a in arts:
            for l in per_article[a['id']]:
                r = rel.setdefault(l['groupId'], {'tabId': l['tabId'], 'tabLabel': l['tabLabel'], 'groupIndex': l['groupIndex'], 'groupId': l['groupId'], 'groupTitle': l['groupTitle'],
                                                  'role': 'related', 'articleIds': [], 'origins': []})
                if l['role'] == 'primary':
                    r['role'] = 'primary'
                if a['id'] not in r['articleIds']:
                    r['articleIds'].append(a['id'])
                if l.get('origin') and l['origin'] not in r['origins']:
                    r['origins'].append(l['origin'])
        groups = sorted(rel.values(), key=lambda g: (0 if g['role'] == 'primary' else 1, order[g['tabId']], g['groupIndex']))
        tabs = collections.OrderedDict()
        for g in groups:
            tabs.setdefault(g['tabId'], {'tabId': g['tabId'], 'label': g['tabLabel'], 'primaryGroupIds': [], 'relatedGroupIds': []})
            tabs[g['tabId']]['primaryGroupIds' if g['role'] == 'primary' else 'relatedGroupIds'].append(g['groupId'])
        tabs_sorted = sorted(tabs.values(), key=lambda t: (0 if t['primaryGroupIds'] else 1, order[t['tabId']]))
        # hero
        h = K.HERO[sd['id']]
        im = imgs[h['id']]
        hero = {'id': h['id'], 'fit': h['fit'], 'alt': im['alt'], 'shows': im['shows'], 'type': im['type'], 'hasAlpha': im['hasAlpha'],
                'srcset': im['srcset'], 'files': [v['file'] for v in im['variants']], 'caveat': h['caveat'],
                'alternates': [{'id': x, 'shows': imgs[x]['shows']} for x in h['alternates']]}
        q = K.INTRO[sd['id']]
        intro = resolve_quote(q, by_key, content)
        intro['verbatim'] = True
        if q.get('note'):
            intro['note'] = q['note']
        intro['alternative'] = resolve_quote(q['alternative'], by_key, content)
        s = {'id': sd['id'], 'slug': sd['slug'], 'name': sd['name'], 'folder': nfc(sd['folder']),
             **({'folderNote': sd['folderNote']} if sd.get('folderNote') else {}),
             'order': len(segs) + 1, 'articleIds': [a['id'] for a in arts], 'articleCount': len(arts),
             'blogCategoryKey': sd['id'], 'priceList': {'tabs': tabs_sorted, 'groups': groups},
             'heroImage': hero, 'intro': intro}
        if sd['id'] == 'dermatologija':
            s['priceList']['note'] = 'On the price list the dermatology services live in the tab "Estetske procedure" (groups "Dermatologija" and "Rast kose"); there is no separate dermatology tab.'
        if sd['id'] == 'telo':
            s['priceList']['note'] = 'The laser epilation article belongs to this segment but its price list is its own tab "Laserska epilacija".'
        if sd['id'] == 'kozmetika':
            s['priceList']['note'] = 'The old site calls this tab "Tretmani lica": all seven articles of the segment are face treatments.'
        segs.append(s)
    return {
        'meta': {'pack': 'Lepa Svaki Dan v2 texts pack: segments.json', 'generatedBy': 'ECHOSYSTEM-MOTHER AGENCY', 'generatedAt': NOW.strftime('%Y-%m-%dT%H:%M:%SZ'),
                 'ids': 'segment.id equals the old blog category key (dermatologija, estetske, kozmetika, telo) and articles[].segmentId; slug is the URL part',
                 'heroImage': 'chosen by what the image really shows (media pack manifest.json alt/shows), not by its file name; fit says how close it is',
                 'intro': 'one verbatim quote from the articles or the old site (content.json); nothing is invented. alternative = a second verbatim sentence.',
                 'counts': {'segments': len(segs), 'articles': sum(s['articleCount'] for s in segs)}},
        'segments': segs,
    }


# ============================================================================ 6. search vocabulary
def fold_map(s):
    s = nfc(s)
    out, mp = [], []
    for i, ch in enumerate(s):
        base = C._TRANSLIT.get(ch)
        if base is None:
            base = ''.join(c for c in unicodedata.normalize('NFD', ch) if unicodedata.category(c) != 'Mn')
        for c in base.lower():
            out.append(c)
            mp.append(i)
    return ''.join(out), mp, s


def find(rx, text):
    """regex on the folded text; returns (matched folded term, snippet of the ORIGINAL text) or None."""
    folded, mp, orig = fold_map(text)
    m = re.search(rx, folded)
    if not m:
        return None
    a, b = mp[m.start()], mp[m.end() - 1] + 1
    snippet = orig[max(0, a - 40):b + 40].strip()
    if len(orig) <= 110:
        snippet = orig
    return m.group(0), snippet


def article_fields(a):
    yield 'title', a['title']
    if a['subtitle']:
        yield 'subtitle', a['subtitle']
    for bi, b in enumerate(a['blocks']):
        if 'text' in b:
            yield 'block %d' % bi, b['text']
        else:
            for ii, it in enumerate(b['items']):
                yield 'block %d item %d' % (bi, ii), it


def build_search(articles, content, per_article, key_to_id):
    order = tab_order(content)
    sk = sid_sort_key(order)
    tabs = {t['id']: t for t in content['pricing']['tabs']}
    svc = {s['id']: s for s in content['allServices']}
    art_by_id = {a['id']: a for a in articles}
    offers = {o['id']: o for o in content['offers']}
    gifts = {p['cenovnik']['id']: i for i, p in enumerate(content['gifts']) if p.get('cenovnik')}
    # service -> [(article id, group link role)] and group -> [(article id, role)]
    svc_links, grp_links = collections.defaultdict(list), collections.defaultdict(list)
    for aid, links in per_article.items():
        for l in links:
            grp_links[l['groupId']].append((aid, l['role'], l['basis']))
            for sid in l['serviceIds']:
                svc_links[sid].append((aid, l['role']))
    art_hit_cache = {}

    def art_evidence(aid, terms):
        k = (aid, tuple(terms))
        if k not in art_hit_cache:
            res = None
            for term in terms:
                for where, text in article_fields(art_by_id[aid]):
                    f = find(term, text)
                    if f:
                        res = {'term': f[0], 'in': where, 'snippet': f[1]}
                        break
                if res:
                    break
            art_hit_cache[k] = res
        return art_hit_cache[k]

    def svc_evidence(sid, terms, entry_articles, gift):
        s = svc[sid]
        t = tabs[sid.split(':')[0]]
        for term in terms:
            for where, text in (('name', s['name']), ('description', s.get('description') or ''), ('group', s['group']), ('tab', t['label'])):
                if text:
                    f = find(term, text)
                    if f:
                        return {'term': f[0], 'in': where, 'snippet': text}
        for aid, role in svc_links.get(sid, []):
            if aid in entry_articles:
                ev = art_evidence(aid, terms)
                if ev:
                    return {'term': ev['term'], 'via': aid, 'in': 'linked article (%s link) %s' % (role, ev['in']), 'snippet': ev['snippet']}
        if gift and sid in gifts:
            return {'term': 'poklon', 'in': 'content.json gifts[%d] (section "Tretmani za poklon")' % gifts[sid], 'snippet': content['sections']['gifts']['intro']}
        return None

    def group_evidence(gid, terms, entry_articles, gift=False):
        tid, gi = gid.split(':')
        g = tabs[tid]['groups'][int(gi)]
        for term in terms:
            f = find(term, g['title'])
            if f:
                return {'term': f[0], 'in': 'group title', 'snippet': g['title']}
        for term in terms:
            for s in g['services']:
                for where, text in (('name', s['name']), ('description', s.get('description') or '')):
                    if text:
                        f = find(term, text)
                        if f:
                            return {'term': f[0], 'in': 'service ' + where + ' (' + s['id'] + ')', 'snippet': text}
        for aid, role, basis in grp_links.get(gid, []):
            if aid in entry_articles:
                ev = art_evidence(aid, terms)
                if ev:
                    return {'term': ev['term'], 'via': aid, 'in': 'linked article (%s link) %s' % (role, ev['in']), 'snippet': ev['snippet']}
        if gift:
            hit = [s['id'] for s in g['services'] if s['id'] in gifts]
            if hit:
                return {'term': 'poklon', 'in': 'content.json gifts[%d] (section "Tretmani za poklon") picks the service %s of this group' % (gifts[hit[0]], hit[0]), 'snippet': content['sections']['gifts']['intro']}
        return None

    def offer_evidence(oid, terms):
        o = offers[oid]
        for term in terms:
            for where, text in (('title', o['title']), ('description', o['description']), ('tag', o['tag'])):
                f = find(term, text)
                if f:
                    return {'term': f[0], 'in': 'offer ' + where, 'snippet': text}
        return None

    def build_entry(e, is_question=False):
        terms = e['terms']
        arts = []
        for k in e['articles']:
            arts.append(key_to_id[k])
        svc_ids = list(e['services'])
        if e['service_scan']:
            sc = e['service_scan']
            rx = re.compile(sc['pattern'])
            for tid in sc['tabs']:
                for g in tabs[tid]['groups']:
                    for s in g['services']:
                        text = s['name'] if sc['field'] == 'name' else s['name'] + ' ' + (s.get('description') or '')
                        if rx.search(C.fold(text)):
                            svc_ids.append(s['id'])
        for gid in e['groups']:
            tid, gi = gid.split(':')
            if tid not in tabs or int(gi) >= len(tabs[tid]['groups']):
                raise SystemExit('entry %s: unknown group %s' % (e['id'], gid))
        for sid in svc_ids:
            if sid not in svc:
                raise SystemExit('entry %s: unknown service %s' % (e['id'], sid))
        for oid in e['offers']:
            if oid not in offers:
                raise SystemExit('entry %s: unknown offer %s' % (e['id'], oid))
        svc_ids = sorted(set(svc_ids), key=sk)
        evidence, problems = {}, []
        for aid in arts:
            if e['direct_evidence']:
                ev = art_evidence(aid, terms)
                if ev:
                    evidence[aid] = ev
                else:
                    problems.append('article ' + aid)
            else:
                evidence[aid] = {'note': 'closest hook: no text supports the phrase directly'}
        for gid in e['groups']:
            ev = group_evidence(gid, terms, set(arts), e['gift_services']) if e['direct_evidence'] else {'note': 'closest hook: no text supports the phrase directly'}
            if ev:
                evidence[gid] = ev
            else:
                problems.append('group ' + gid)
        for sid in svc_ids:
            ev = svc_evidence(sid, terms, set(arts), e['gift_services']) if e['direct_evidence'] else {'note': 'closest hook: no text supports the phrase directly'}
            if ev:
                evidence[sid] = ev
            else:
                problems.append('service ' + sid)
        for oid in e['offers']:
            ev = offer_evidence(oid, terms)
            if ev:
                evidence[oid] = ev
            else:
                problems.append('offer ' + oid)
        if problems:
            raise SystemExit('search entry %r: no textual support for %s (terms %s)' % (e['id'], ', '.join(problems), terms))
        seg_ids = []
        for aid in arts:
            sid_ = art_by_id[aid]['segmentId']
            if sid_ not in seg_ids:
                seg_ids.append(sid_)
        caution = None
        if e['caution']:
            c_arts = [key_to_id[k] for k in e['caution']['articles']]
            for aid in c_arts:
                if not art_evidence(aid, terms):
                    raise SystemExit('caution article %s of %s has no mention of the topic' % (aid, e['id']))
            caution = {'articleIds': c_arts, 'note': e['caution']['note']}
        # phrases: lowercase, with and without diacritics
        phrases = []
        for p in e['phrases']:
            for v in (nfc(p).lower(), C.fold(p)):
                if v not in phrases:
                    phrases.append(v)
        out = {'id': e['id']}
        if is_question:
            out['question'] = e['phrases'][0][:1].upper() + e['phrases'][0][1:] + ('' if e['phrases'][0].endswith('?') else '?')
        out.update({'phrases': phrases,
                    'targets': {'groupIds': list(e['groups']), 'serviceIds': svc_ids, 'articleIds': arts, 'segmentIds': seg_ids, 'offerIds': list(e['offers'])},
                    'strength': e['strength'], 'note': e['note'], 'evidence': evidence})
        if caution:
            out['caution'] = caution
        if not e['direct_evidence']:
            out['directEvidence'] = False
        return out

    entries = [build_entry(e) for e in SD.ENTRIES]
    questions = [build_entry(e, True) for e in SD.QUESTIONS]
    assert len(questions) == 12, len(questions)
    # coverage of the phrases the brief requires
    have = {C.fold(p) for e in entries for p in e['phrases']}
    missing = [r for r in SD.REQUIRED if r not in have]
    if missing:
        raise SystemExit('required phrases missing from the vocabulary: %s' % missing)
    # phrase index (folded phrase -> entry ids)
    idx = collections.OrderedDict()
    for e in entries + questions:
        for p in e['phrases']:
            f = C.fold(p)
            idx.setdefault(f, [])
            if e['id'] not in idx[f]:
                idx[f].append(e['id'])
    ids = [e['id'] for e in entries + questions]
    assert len(ids) == len(set(ids)), 'duplicate entry ids'
    return {
        'meta': {
            'pack': 'Lepa Svaki Dan v2 texts pack: search-synonyms.json', 'generatedBy': 'ECHOSYSTEM-MOTHER AGENCY', 'generatedAt': NOW.strftime('%Y-%m-%dT%H:%M:%SZ'),
            'display': 'HIDDEN vocabulary: use it to find results, never render the phrases as hashtags, chips or tags.',
            'normalization': 'Fold both the query and the phrases before comparing: lowercase, c/c/s/z for c-caron/c-acute/s-caron/z-caron, "dj" for d-stroke, drop combining marks. Every phrase is already stored with and without diacritics; phraseIndex is keyed by the folded form.',
            'targets': {'groupIds': '"<tab>:<groupIndex>" (service-links.json)', 'serviceIds': 'content.json allServices[].id', 'articleIds': 'art-1 ... art-21',
                        'segmentIds': 'segments.json (derived from the targeted articles)', 'offerIds': 'content.json offers[].id (extra key: the old site searches offers too)'},
            'strength': {'strong': 'the target is the subject of the phrase', 'medium': 'the text supports it, it is not the main subject or the wording differs',
                         'weak': 'closest hook only, show as a soft suggestion and never as a treatment claim'},
            'honesty': 'Every target contains one of the entry terms in its own name/description/title/text, or is linked to a targeted article that does (see evidence). Entries with directEvidence:false are the only exception and say so.',
            'counts': {'entries': len(entries), 'questions': len(questions), 'phrases': len(idx),
                       'requiredPhrasesCovered': len(SD.REQUIRED), 'weakEntries': [e['id'] for e in entries if e['strength'] == 'weak']},
        },
        'entries': entries,
        'questions': questions,
        'phraseIndex': idx,
    }


# ============================================================================ 7. CHANGES.md / MANIFEST-ARTICLES.md
def diff_strings(old, new, path=''):
    """every string that differs between two content trees (same shape): [(path, before, after)]."""
    out = []
    if isinstance(old, dict) and isinstance(new, dict):
        for k in old:
            if k in new:
                out += diff_strings(old[k], new[k], path + ('.' if path else '') + k)
    elif isinstance(old, list) and isinstance(new, list):
        for i in range(min(len(old), len(new))):
            out += diff_strings(old[i], new[i], '%s[%d]' % (path, i))
    elif isinstance(old, str) and isinstance(new, str) and old != new:
        out.append((path, old, new))
    return out


def content_changes(old, new):
    """authoritative list of what changed in content.json (backup vs now), classified."""
    rows = []
    for path, b, a in diff_strings(old, new):
        if path.startswith('_meta') and not (path.startswith('_meta.warnings') and (EM in b or EN in b)):
            continue
        hits = []
        explained = C.fix_text(b, hits) == a
        if C.OWNER_RE.search(b) or b in C.RULES['owner']['exactStrings']:
            kind = 'owner'
        elif EM in b or EN in b:
            kind = 'dash'
        elif re.fullmatch(r'0\d:\d\d', b):
            kind = 'numbering'
        else:
            kind = 'other'
        rows.append({'path': path, 'before': b, 'after': a, 'kind': kind, 'rules': hits, 'explained': explained})
    return rows


def md_cell(s):
    return str(s).replace('|', '\\|').replace('\n', ' ')


def window(before, after, width=70):
    """the part of two strings around their first difference (so long strings still show what changed)."""
    b, a = C.printable(before), C.printable(after)
    p = 0
    while p < min(len(b), len(a)) and b[p] == a[p]:
        p += 1
    start = max(0, p - 35)
    pre = '...' if start else ''
    return pre + b[start:start + width], pre + a[start:start + width]


def render_changes(content_log, art_log, doc_stats, content, today):
    L = []
    P = L.append
    old_content = C.read_json(os.path.join(C.BACKUP_DIR, 'content.json'))
    rows = content_changes(old_content, content)
    unexplained = [r for r in rows if r['kind'] in ('owner', 'dash') and not r['explained']]
    if unexplained:
        raise SystemExit('content changes not explained by the reviewed rules: %s' % [r['path'] for r in unexplained])
    owner = [r for r in rows if r['kind'] == 'owner']
    dash_c = [r for r in rows if r['kind'] == 'dash']
    other_c = [r for r in rows if r['kind'] == 'other']
    num_c = [r for r in rows if r['kind'] == 'numbering']
    dash_a = [c for c in art_log if any(r.startswith('dash') for r in c['rules'])]
    owner_a = [c for c in art_log if any(r.startswith('owner') for r in c['rules'])]
    P('# CHANGES: shared content corrections')
    P('')
    P('Applied to `content.json` and `content.js` (regenerated by `_tools/build-content.js`, which now applies the fixes itself) and to the new texts pack (`articles.json`, `segments.json`, `service-links.json`, `search-synonyms.json`). Rules: `_tools/content-fixes.rules.json` (one set, read by the JS and the Python tools).')
    P('')
    P('The dash characters themselves are never written in this file: `{U+2014}` is an em dash, `{U+2013}` an en dash.')
    P('')
    P('| Fix | content.json / content.js | articles (Word documents) | search / segments / links |')
    P('|---|---|---|---|')
    P('| (a) owner is not a doctor | %d strings (+ new field `founder.signature`) | %d | 0 |' % (len(owner), len(owner_a)))
    P('| (b) em dash removed, other non-numeric dashes replaced | %d strings | %d strings | 0 (written clean) |' % (len(dash_c), len(dash_a)))
    P('| (c) leading zeros in numbering | %d derived hour strings ("09:00" -> "9:00") | 0 | 0 (ids are art-1 ... art-21, no zero padding) |' % len(num_c))
    P('| prices, service ids, counts | unchanged (checksums below) | n/a | n/a |')
    P('')
    P('## A. The owner is not a doctor')
    P('')
    P('Rule: remove "Dr"/"dr" directly before "Tamara" (regex `\\b[Dd]r\\.?\\s+(?=Tamara\\b)`), add `founder.signature` = "%s" and use it wherever the owner is signed (`founder.signatureName`).' % C.SIGNATURE)
    P('')
    P('| Where | Before | After |')
    P('|---|---|---|')
    for c in owner:
        P('| `%s` | %s | %s |' % (c['path'], md_cell(C.printable(c['before'])), md_cell(C.printable(c['after']))))
    P('| `founder.signature` | (field did not exist) | %s |' % md_cell(C.SIGNATURE))
    for c in owner_a:
        P('| %s, %s | %s | %s |' % (c['file'], c['where'], md_cell(C.printable(c['before'])), md_cell(C.printable(c['after']))))
    P('')
    P('Not changed on purpose (they name other people, or only a surname, and must not be assumed to be the owner):')
    P('')
    P('- `Specijalistički dermatološki pregled dr Krstić` (price list, estetske:5:0)')
    P('- `Hijaluron 1ml, dr spec. Vujačić` (price list, estetske:1:2)')
    P('- the "dr" in front of the blanks in the Word documents: it belongs to another doctor, whose name is still missing (see placeholders in `articles.json`)')
    P('')
    P('Needs the client: after the change the price-list line `Specijalistički pregled Tamara Vujačić` (estetske:0:1) reads as a specialist exam by a person who is not a doctor.')
    P('')
    P('## B. Dashes')
    P('')
    P('Rule: no em dash (U+2014) anywhere; an en dash (U+2013) only between numbers (`400{U+2013}500€`, `2.800{U+2013}4.700 rsd`). Each replacement is a reviewed, per-sentence rule (comma, colon, plain hyphen); an unknown dash stops the build instead of being guessed.')
    P('')
    P('### B1. content.json / content.js')
    P('')
    P('| Where | Before | After | Replaced by |')
    P('|---|---|---|---|')
    rule_why = {r['find']: r['why'] for r in C.RULES['dash']['rules']}
    for c in dash_c:
        why = []
        for r in c['rules']:
            if r.startswith('dash:placeholder'):
                why.append(C.RULES['dash']['exactWhy'])
            elif r.startswith('dash:'):
                key = r[5:].replace('{U+2014}', EM).replace('{U+2013}', EN)
                why.append(rule_why.get(key, ''))
        wb, wa = window(c['before'], c['after'])
        P('| `%s` | %s | %s | %s |' % (c['path'], md_cell(wb), md_cell(wa), md_cell('; '.join(dict.fromkeys(why)))))
    if other_c:
        raise SystemExit('unclassified content change: %s' % [r['path'] for r in other_c])
    P('')
    P('Left as they are (numeric ranges, allowed): %d en dashes in price strings such as `400{U+2013}500€ / 46.600{U+2013}58.600 rsd` and `od 24{U+2013}40€`.' % sum(1 for _ in re.finditer(EN, json.dumps(content, ensure_ascii=False))))
    P('')
    P('The `blog[].anchor` / `anchorByIndex` / `postAnchor` values are percent-encoded keys of the OLD blog.html posts; they are not text and were not changed (one of them still encodes the old en dash as `%E2%80%93`).')
    P('')
    P('### B2. Word documents (articles.json)')
    P('')
    P('| Document | Where | Before | After | Replaced by |')
    P('|---|---|---|---|---|')
    for c in dash_a:
        why = []
        for r in c['rules']:
            if r.startswith('dash:'):
                key = r[5:].replace('{U+2014}', EM).replace('{U+2013}', EN)
                why.append(rule_why.get(key, ''))
        wb, wa = window(c['before'], c['after'])
        P('| %s | %s | %s | %s | %s |' % (md_cell(os.path.basename(c['file'])), c['where'], md_cell(wb), md_cell(wa), md_cell('; '.join(dict.fromkeys(why)))))
    P('')
    P('## C. Numbering without leading zeros')
    P('')
    P('- No zero-padded counters existed in the data (the old site has none in the extracted strings).')
    P('- The only zero-padded numbers were derived by the generator: `hero.bar.hoursParsed.open` and `contact.hours.parsed.open` were "09:00"; now "9:00" (`close` stays "21:00"). The generator now strips the zero.')
    P('- New data uses plain numbers: article ids `art-1` ... `art-21`, segment `order` 1 ... 4, tables in the manifests start at 1.')
    P('- Not numbering, so untouched: the price-list date `15.03.2026.` (shown verbatim by the old site), `pricing.validFromISO` (an ISO date) and the generation timestamps.')
    P('')
    P('## D. Other changes made while doing this')
    P('')
    P('- `content.json` `_meta.checksums`: `serviceLinesSha256` keeps its old value (it hashes the SOURCE lines and is therefore unchanged); new `serviceLinesSha256AfterFixes` (one service name changed: estetske:0:1) and the name-independent `priceLinesSha256` (id|price|priceEur|priceRsd, identical before and after the fixes). `_meta.fixes` summarises this file.')
    P('- `content.json` `_meta.warnings`: +4 (owner title in a price-list line, two doctor titles left in place, founder photo alt text) and the "Preime" note now says the Word document spells it the same way in title, body and source URL.')
    P('- `_tools/build-content.js`, `verify-content.js`, `mutation-test.js`, `write-manifest.js`: em dashes removed from comments and templates (logic literals are now `\\u2014` escapes), fixes applied in memory before counts/checksums, the verifier passes the old site text through the same reviewed fixes, new checks (owner title, dashes, checksums), `fixes-log.content.json` is written, every generated text is asserted dash-clean.')
    P('- `COPY_NOTES.md`, `MANIFEST.md`, `_tools/verification-report.json`: regenerated; no em dash left in any of them.')
    P('- Document text normalisation (not content changes): runs of spaces/tabs collapsed and trailing spaces trimmed in every paragraph (%d non-breaking spaces turned into normal spaces); %d empty paragraphs skipped; the title line is split at the first colon.' % (doc_stats['nbsp'], doc_stats['emptyParagraphs']))
    P('- The blanks `____` in six sentences became `[[dopuniti: ime lekara]]` (5) and `[[dopuniti: naziv lasera]]` (1). No name was invented.')
    P('')
    P('## E. Not touched')
    P('')
    P('- The three previews (`a-lumine`, `b-alder`, `c-luzia`) and `_deploy/` keep their own slim copies of the content (and a few strings in their HTML): they still carry the old owner title and em dashes until they are re-synced from this pack.')
    P('- Typos in the Word documents are kept verbatim (listed in `MANIFEST-ARTICLES.md`).')
    P('')
    P('---')
    P('Pripremio ECHOSYSTEM-MOTHER AGENCY · ' + today)
    P('')
    return '\n'.join(L)


def render_manifest(articles, segments, links, search, content, report, vreport, today, n_links, n_hints):
    L = []
    P = L.append
    seg_name = {s['id']: s['name'] for s in segments['segments']}
    P('# MANIFEST-ARTICLES: Lepa Svaki Dan v2 texts pack')
    P('')
    P('The treatment texts of the 21 Word documents (written by the owner, Tamara, with a copywriter) as data for the three v2 previews, plus the price-list links and a hidden search vocabulary. Complements `MANIFEST.md` (the old-site content pack). Generated %s by ECHOSYSTEM-MOTHER AGENCY with `python3 _tools/build-texts-pack.py`.' % NOW.strftime('%Y-%m-%dT%H:%M:%SZ'))
    P('')
    P('- Verification: **%s** (%d checks, `python3 _tools/verify-texts-pack.py`).' % ('ALL PASSED' if vreport['ok'] else 'FAILED', len(vreport['checks'])))
    P('- Language: Serbian (sr), diacritics preserved. Reviewed content fixes are applied (see `CHANGES.md`): the owner is not a doctor, no em dashes.')
    P('')
    P('## 1. Files (all in `v2/_shared/content/`)')
    P('')
    P('| File | What it is |')
    P('|---|---|')
    P('| `articles.json` | Array of **%d** article objects (see section 2). UTF-8 JSON. |' % len(articles))
    P('| `articles.js` | `window.LSD_ARTICLES = [...];` identical data, loads from `file://` via `<script src="../_shared/content/articles.js">`. |')
    P('| `segments.json` | **4** segments: ordered article ids, price-list tabs/groups, suggested hero image, one verbatim intro quote. |')
    P('| `service-links.json` | Per article the related price-list groups and services (primary/related, with the basis of every match), the two old SERVICE_BLOG_LINKS rules reconciled, unmatched articles (none), reverse index group -> articles and service -> articles. |')
    P('| `search-synonyms.json` | HIDDEN search vocabulary: **%d** entries + **%d** invented natural-language questions, phrase variants with and without diacritics -> groups/services/articles/segments/offers, with a note and machine-checked evidence per target. |' % (len(search['entries']), len(search['questions'])))
    P('| `CHANGES.md` | Every replacement made by the content fixes. |')
    P('| `MANIFEST-ARTICLES.md` | This file. |')
    P('| `_tools/build-texts-pack.py`, `_tools/texts_curation.py`, `_tools/texts_search_data.py`, `_tools/lsd_common.py` | Builder and the reviewed curation data (price-list links, search entries, hero images, comment resolutions). |')
    P('| `_tools/verify-texts-pack.py`, `_tools/verification-report.texts.json` | Independent verifier and its machine-readable report. |')
    P('| `_tools/content-fixes.js`, `_tools/content-fixes.rules.json`, `_tools/fixes-log.content.json`, `_tools/fixes-log.articles.json` | The shared fix rules and the logs of what they changed. |')
    P('')
    P('Source documents: `/Users/vladajovic/Developer/LEPA SVAKI DAN/procedure/` (4 folders = 4 segments; the folder "Estske procedure" is spelled with a typo, the segment is "Estetske procedure").')
    P('')
    P('## 2. Shape of `articles.json`')
    P('')
    P('```json')
    P('{ "id": "art-6", "slug": "botulinski-toksin-tipa-a-botoks-tretman", "segmentId": "estetske",')
    P('  "title": "...", "subtitle": "", "blocks": [ {"type":"h2","text":"..."}, {"type":"p","text":"...","links":[...]}, {"type":"ul","items":["..."]} ],')
    P('  "excerpt": "...", "readMinutes": 3, "sourceFile": "procedure/.../x.docx",')
    a6 = next(a for a in articles if a['id'] == 'art-6')
    P('  "hasPlaceholders": false, "placeholders": [], "wordCount": %d, "titleWordCount": %d,' % (a6['wordCount'], a6['titleWordCount']))
    P('  "legacy": {...}, "crossLinkHints": [...] }')
    P('```')
    P('')
    P('| Field | Meaning |')
    P('|---|---|')
    P('| `id`, `slug`, `segmentId` | `art-1` ... `art-21` in document order (= old blog order); `slug` = ASCII kebab-case of the title; `segmentId` = `dermatologija`, `estetske`, `kozmetika` or `telo` (the old blog category keys). |')
    P('| `title`, `subtitle` | The first line of the document split at its FIRST colon. `subtitle` is `""` when the line has no colon (%d of 21 first lines have a colon, %d do not). |' % (sum(1 for a in articles if a['subtitle']), sum(1 for a in articles if not a['subtitle'])))
    P('| `blocks[]` | Document order and structure preserved: `h2` (Word Heading 2), `h3` (Heading 3, **used in %d articles; a fifth type beyond h2/p/ul/ol, because the documents have sub-sections**), `p`, `ul` / `ol` with `items[]`. `ol` only where Word numbering is decimal (%d list); every other list is a bullet list. Empty paragraphs are skipped. The title line is not a block. |' % (sum(1 for a in articles if any(b['type'] == 'h3' for b in a['blocks'])), sum(1 for a in articles for b in a['blocks'] if b['type'] == 'ol')))
    P('| `blocks[].links[]` | Hyperlinks the copywriter put in the Word text (%d): `{text, legacyHref, target}` (lists also carry `item`). `legacyHref` is the URL on the live lepasvakidan.rs site, kept for reference only. `target` = `{type: article\\|segment\\|pricelist\\|external\\|none, id}` says where it should point in the new site. |' % n_links)
    P('| `blocks[].flag` | Only on one paragraph (Preime Dermafacial): a copywriter note naming the site the text was taken from, not article copy. |')
    P('| `excerpt` | First 170 characters of the first paragraph, cut at a word boundary, no ellipsis; a trailing colon/comma is dropped. |')
    P('| `readMinutes`, `wordCount` | `wordCount` = words of all blocks (a word = a token with a letter or digit; every placeholder counts as one word; the title is not included); `readMinutes` = ceil(wordCount / 200), minimum 1. |')
    P('| `placeholders[]` | The blanks of the Word documents: `{token, kind, blockIndex, blockType, itemIndex?, context}`. The blank is replaced in the text by `[[dopuniti: ime lekara]]` (a doctor) or `[[dopuniti: naziv lasera]]` (a laser device), chosen from the sentence. The "dr" in front of a blank belongs to another doctor and is left in place. **Never show a text with a token on a public page.** |')
    P('| `legacy` | The matching old blog post in `content.json` `blog[]`: `blogIndex` (= article number - 1), `blogTitle`, `titleMatch` (`exact`, or `order` when the title changed), `blogWordCount`, `blogHadBlanks`. |')
    P('| `crossLinkHints[]` | The copywriter\'s Word comments (%d): which words should link to which other text. `{comment, anchorText, blockIndex, resolution: confident\\|partial\\|none, basis, targetArticleIds, targetSegmentIds}`. Hints are data, not instructions to the builders. |' % n_hints)
    P('')
    P('## 3. Counts')
    P('')
    seg_counts = collections.Counter(a['segmentId'] for a in articles)
    P('| Item | Count |')
    P('|---|---|')
    P('| Articles | **%d** (%s) |' % (len(articles), ', '.join('%s %d' % (seg_name[s['id']], seg_counts[s['id']]) for s in segments['segments'])))
    P('| Blocks | %d (h2 %d, h3 %d, p %d, ul %d, ol %d) |' % (sum(len(a['blocks']) for a in articles), *[sum(1 for a in articles for b in a['blocks'] if b['type'] == t) for t in ('h2', 'h3', 'p', 'ul', 'ol')]))
    P('| Words (all articles) | %s |' % '{:,}'.format(sum(a['wordCount'] for a in articles)))
    P('| Articles with placeholders | **%d** (%d blanks) |' % (sum(1 for a in articles if a['hasPlaceholders']), sum(len(a['placeholders']) for a in articles)))
    P('| Hyperlinks in the texts / Word comments | %d / %d |' % (n_links, n_hints))
    P('| Price-list groups linked / without an article | %d of %d / %d |' % (links['meta']['counts']['groups'] - links['meta']['counts']['groupsWithoutArticle'], links['meta']['counts']['groups'], links['meta']['counts']['groupsWithoutArticle']))
    P('| Services linked to an article | %d of %d |' % (links['meta']['counts']['servicesLinked'], links['meta']['counts']['services']))
    P('| Unmatched articles | %d |' % links['meta']['counts']['unmatchedArticles'])
    P('| Search entries / questions / phrases | %d / %d / %d |' % (len(search['entries']), len(search['questions']), len(search['phraseIndex'])))
    P('')
    P('## 4. Article -> segment -> price-list groups -> placeholders')
    P('')
    P('Groups: **bold** = primary (the group is the subject), plain = related. Group ids are `<tab>:<index>`.')
    P('')
    P('| # | Article | Segment | Words | Min | Price-list groups | Placeholders |')
    P('|---|---|---|---|---|---|---|')
    by_art = {x['articleId']: x for x in links['articles']}
    for a in articles:
        la = by_art[a['id']]
        gl = []
        for g in la['groups']:
            label = '%s `%s` %s' % (g['groupTitle'], g['groupId'], '')
            gl.append(('**%s** `%s`' if g['role'] == 'primary' else '%s `%s`') % (g['groupTitle'], g['groupId']))
        ph = '%d (%s)' % (len(a['placeholders']), ', '.join(sorted({p['kind'] for p in a['placeholders']}))) if a['placeholders'] else 'none'
        P('| %s | %s | %s | %d | %d | %s | %s |' % (a['id'].replace('art-', ''), md_cell((a['title'] + (': ' + a['subtitle'] if a['subtitle'] else ''))[:95]), seg_name[a['segmentId']], a['wordCount'], a['readMinutes'], '; '.join(gl), ph))
    P('')
    P('Articles WITHOUT a price-list group: **none**. Price-list groups without an article: %s.' % ', '.join('%s `%s`' % (g['title'], g['groupId']) for g in links['groupsWithoutArticle']))
    P('')
    P('## 5. Open items for the client')
    P('')
    ph_arts = [a for a in articles if a['hasPlaceholders']]
    P('1. **Missing names.** %d blanks in %d articles: %s. The doctor\'s name (and the Candela laser model available in the centre) must come from the client. Nothing was invented.' % (sum(len(a['placeholders']) for a in ph_arts), len(ph_arts), '; '.join('%s (%d)' % (a['title'], len(a['placeholders'])) for a in ph_arts)))
    P('2. **Possible hint, not used:** the price list names "dr Krstić" for the specialist dermatological exam (estetske:5:0). It may be the doctor of the blanks, but the texts do not say so; confirm before filling anything in.')
    P('3. `Specijalistički pregled Tamara Vujačić` (estetske:0:1): see `CHANGES.md` section A.')
    P('4. The first line of `Klasična dermatologija.docx` reads "Dermatološki/pregled dermatologa" (no space after the slash). Kept verbatim; the old blog title was "Klasična dermatologija: Šta obuhvata dermatološki pregled?" (`legacy.blogTitle`).')
    P('5. The title "Bezbolna i bezbedna: Laserska epilacija Venus Velocity aparatom" gives the slug `bezbolna-i-bezbedna` (title = part before the colon, as specified). Consider the slug `laserska-epilacija-venus-velocity` when the real URLs are decided.')
    P('6. The copywriter links point to the live site (lepasvakidan.rs, see `links[].legacyHref`); `links[].target` maps them to this pack where a text exists. No text exists for mesotherapy, massages, the "Glow&Go" protocols or a doctor biography.')
    P('7. Preime Dermafacial: the Word text ends with a note naming the manufacturer\'s website as the source of the text (`blocks[].flag`). Check the rights before publishing that text.')
    P('8. Typos kept verbatim for the copywriter to confirm (art-N = article number): "leserske" (art-1), "pigmentih" (art-1), "interdermalni" (art-3), "stpljivim" and "pišete" (art-4), "odgovori" for "odgovorni" (art-6), "liftiga", "majči pogled" and "izgleda očiju" (art-8), "prvu prvenciju" (art-9), "boda" for "voda" and "važa koža" (art-14), "rezulate", "tekturu" and "eletromagnetna" (art-15), "kolegana" (art-16), "Najčešćše" (art-17), "čuščje" and "čineću" (art-18; the price list says "guščije"), "elektromangnetna" (art-20), "laserke epilacije" (art-21); "Miolift" next to "MyoLift" (art-12).')
    P('9. The three previews (a-lumine, b-alder, c-luzia) and `_deploy/` still carry the old owner title and em dashes in their own content copies; they need a re-sync from this pack (see `CHANGES.md` section E).')
    P('')
    P('## 6. Verification (`_tools/verify-texts-pack.py`)')
    P('')
    for c in vreport['checks']:
        P('- %s %s' % ('PASS' if c['ok'] else 'FAIL', c['name']))
    P('')
    P('Word counts per document (independent recounts; "source" = every paragraph of the Word file counted straight from the XML, python-docx = its own text, textutil = macOS converter, pack = title + blocks of the pack; a blank or placeholder counts as one word, a lone dash as none):')
    P('')
    P('| # | Document | source (XML) | python-docx | textutil | pack | old blog |')
    P('|---|---|---|---|---|---|---|')
    for r in vreport['wordTable']:
        P('| %s | %s | %s | %s | %s | %s | %s |' % (r['n'], md_cell(r['file']), r['xml'], r['pydocx'], r['textutil'], r['pack'], r['legacy']))
    P('')
    P('## 7. Rebuild')
    P('')
    P('```bash')
    P('cd "/Users/vladajovic/Developer/LEPA SVAKI DAN/v2/_shared/content"')
    P('node _tools/build-content.js          # content.json/js (+ fixes), MANIFEST.md, COPY_NOTES.md')
    P('python3 _tools/build-texts-pack.py    # articles/segments/links/search + CHANGES.md + this file (pip3 install --user python-docx)')
    P('python3 _tools/verify-texts-pack.py   # exit code 1 on any failed check')
    P('```')
    P('')
    P('---')
    P('Pripremio ECHOSYSTEM-MOTHER AGENCY · ' + today)
    P('')
    return '\n'.join(L)


# ============================================================================ main
def load_verifier():
    spec = importlib.util.spec_from_file_location('verify_texts_pack', os.path.join(HERE, 'verify-texts-pack.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def main():
    content = C.read_json(os.path.join(C.CONTENT_DIR, 'content.json'))
    img_manifest = C.read_json(os.path.join(C.IMG_DIR, 'manifest.json'))
    docs = list_documents()
    if len(docs) != 21:
        raise SystemExit('expected 21 documents, found %d' % len(docs))
    report = {}
    fixlog = FixLog()
    articles = build_articles(docs, content, fixlog, report)
    key_to_id = {a['_key']: a['id'] for a in articles}
    pub_articles, n_links, n_hints = finish_articles(articles, key_to_id, None)

    per_article = build_service_links(articles, content, key_to_id)
    old_rules = reconcile_old_rules(content, per_article, key_to_id)
    links = render_service_links(pub_articles, content, per_article, old_rules)
    segments = build_segments(pub_articles, content, per_article, img_manifest)
    search = build_search(pub_articles, content, per_article, key_to_id)
    for p in pub_articles:
        p.pop('_key', None)

    # ---------------------------------------------------------------- write data files
    NOWS = NOW.strftime('%Y-%m-%dT%H:%M:%SZ')
    out_dir = C.CONTENT_DIR
    arts_json = C.json_no_dash(pub_articles, 2)
    C.assert_clean(arts_json, 'articles.json')
    C.write_text(os.path.join(out_dir, 'articles.json'), arts_json + '\n')
    arts_js = ('/* Lepa Svaki Dan v2 texts pack: 21 articles from the Word documents, generated by _tools/build-texts-pack.py (ECHOSYSTEM-MOTHER AGENCY). Works from file:// */\n'
               'window.LSD_ARTICLES = ' + arts_json.replace(chr(0x2028), '\\u2028').replace(chr(0x2029), '\\u2029') + ';\n')
    C.assert_clean(arts_js, 'articles.js')
    C.write_text(os.path.join(out_dir, 'articles.js'), arts_js)
    for name, obj in (('segments.json', segments), ('service-links.json', links), ('search-synonyms.json', search)):
        txt = C.json_no_dash(obj, 2)
        C.assert_clean(txt, name)
        C.write_text(os.path.join(out_dir, name), txt + '\n')
    # logs
    art_log = fixlog.items
    C.write_text(os.path.join(HERE, 'fixes-log.articles.json'), C.json_no_dash({'generatedAt': NOWS, 'changes': art_log}, 2) + '\n')
    content_log = C.read_json(os.path.join(HERE, 'fixes-log.content.json'))['changes']
    changes_md = render_changes(content_log, art_log, report['docStats'], content, TODAY)
    C.assert_clean(changes_md, 'CHANGES.md')
    C.write_text(os.path.join(out_dir, 'CHANGES.md'), changes_md)

    # ---------------------------------------------------------------- verify, then the manifest (embeds the report)
    V = load_verifier()
    vreport = V.run()
    C.write_text(os.path.join(HERE, 'verification-report.texts.json'), C.json_no_dash(vreport, 2) + '\n')
    manifest = render_manifest(pub_articles, segments, links, search, content, report, vreport, TODAY, n_links, n_hints)
    C.assert_clean(manifest, 'MANIFEST-ARTICLES.md')
    C.write_text(os.path.join(out_dir, 'MANIFEST-ARTICLES.md'), manifest)

    # console summary
    print('articles: %d | segments: %d | blocks: %d | words: %d | placeholders: %d in %d articles'
          % (len(pub_articles), len(segments['segments']), sum(len(a['blocks']) for a in pub_articles), sum(a['wordCount'] for a in pub_articles),
             sum(len(a['placeholders']) for a in pub_articles), sum(1 for a in pub_articles if a['hasPlaceholders'])))
    print('links: %d hyperlinks, %d comments | pseudo-heading candidates: %s' % (n_links, n_hints, report['docStats']['pseudoHeadingCandidates'] or 'none'))
    print('search: %d entries, %d questions, %d phrases' % (len(search['entries']), len(search['questions']), len(search['phraseIndex'])))
    print('verification: %s (%d checks)' % ('OK' if vreport['ok'] else 'FAILED', len(vreport['checks'])))
    print()
    print('%-4s %-46s %6s %6s %6s %6s %6s' % ('#', 'document', 'xml', 'pydocx', 'textut', 'pack', 'oldblg'))
    for r in vreport['wordTable']:
        print('%-4s %-46s %6s %6s %6s %6s %6s' % (r['n'], r['file'][:46], r['xml'], r['pydocx'], r['textutil'], r['pack'], r['legacy']))
    if not vreport['ok']:
        for c in vreport['checks']:
            if not c['ok']:
                print('FAIL', c['name'], c.get('detail', ''))
        sys.exit(1)


if __name__ == '__main__':
    main()
